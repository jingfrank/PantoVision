#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
弓网多算法分层检测体系 - 典型视频离线基准评测套件 (test_benchmark_suite.py)
对 ./弓网典型视频/ 或客户误报视频目录下的所有工况视频运行统一检测调度层与融合引擎，
验证 6 大类故障（火花、脏污、倾斜、飘弓、异物、结构异常）的识别准确率、误报抑制率与推理时延。
"""

import os
import sys
import time
import argparse
import glob
import json
import csv
import pathlib
from collections import defaultdict
import cv2
import numpy as np
from typing import Dict, List, Any, Optional

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)
FASTAD_CORE = os.path.join(BASE_DIR, "fastad_core")
if os.path.exists(FASTAD_CORE) and FASTAD_CORE not in sys.path:
    sys.path.append(FASTAD_CORE)

from core.data_types import FramePacket, FusionResult, DetectionContext
from core.scheduler import DetectionScheduler
from core.fusion import FusionEngine, FAULT_NAME_MAP
from core.output import Visualizer

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}
VIDEO_EXTS = {".mp4", ".avi", ".mkv", ".mov", ".flv"}

FASTAD_FAULT_TYPES = {"dirty", "foreign_object", "structure_abnormal"}

FAULT_CN_MAP_SIMPLE = {
    "dirty": "画面脏污",
    "foreign_object": "悬挂异物",
    "structure_abnormal": "受电弓结构异常",
}


def _render_fastad_diag_img(
    frame: np.ndarray,
    fastad_result: dict,
    fault_type: str,
    score: float,
    alpha: float = 0.40,
    vmax: float = 4.0,
) -> np.ndarray:
    """
    [仅测试与离线评测启用] 将 FastAD 推理结果中的 anomaly_map 与原帧透明叠加，
    并拼接顶部诊断 header 条与缺陷候选框。

    注意：anomaly_map 已由 PantographFastADModel.predict() 内部 resize
    至原图分辨率，此处无需再次 resize（仅作保险尺寸检查）。
    """
    h, w = frame.shape[:2]
    anomaly_map = fastad_result.get("anomaly_map")
    if anomaly_map is None:
        return frame

    bboxes = fastad_result.get("bboxes", [])
    threshold = fastad_result.get("image_threshold", 1.0)

    # 保险尺寸检查
    if anomaly_map.shape[:2] != (h, w):
        anomaly_map = cv2.resize(anomaly_map, (w, h), interpolation=cv2.INTER_LINEAR)

    # 伪彩色热图透明混合 (与 model.generate_overlay 逻辑保持一致)
    norm_map = np.clip(anomaly_map / (vmax + 1e-8), 0.0, 1.0)
    heatmap_u8 = (norm_map * 255.0).astype(np.uint8)
    colored_heatmap = cv2.applyColorMap(heatmap_u8, cv2.COLORMAP_JET)
    blended = cv2.addWeighted(frame, 1.0 - alpha, colored_heatmap, alpha, 0)

    # 绘制缺陷候选框（红色，由 binary_mask 轮廓提取）
    for b in bboxes:
        bx1, by1, bx2, by2 = [int(v) for v in b]
        cv2.rectangle(blended, (bx1, by1), (bx2, by2), (0, 0, 255), 3, cv2.LINE_AA)

    # 顶部工业诊断条（深灰背景，青色文字）
    header_h = 52
    header = np.zeros((header_h, w, 3), dtype=np.uint8)
    header[:] = (30, 30, 30)
    cn_name = FAULT_CN_MAP_SIMPLE.get(fault_type, FAULT_NAME_MAP.get(fault_type, fault_type))
    label = f"FastAD Diag | {cn_name} | Score: {score:.3f} (Thr: {threshold:.3f})"
    cv2.putText(
        header, label, (16, 34), cv2.FONT_HERSHEY_SIMPLEX, 0.82,
        (0, 255, 255), 2, cv2.LINE_AA
    )

    return np.vstack([header, blended])


def collect_videos_recursive(base_dir: str) -> List[str]:
    """递归扫描目录，收集所有视频文件（大小写扩展名兼容）"""
    root = pathlib.Path(base_dir)
    return sorted(
        str(p) for p in root.rglob("*")
        if p.is_file() and p.suffix.lower() in VIDEO_EXTS
    )


def evaluate_video(
    video_path: str,
    scheduler: DetectionScheduler,
    fusion: FusionEngine,
    visualizer: Visualizer,
    output_dir: str,
    start_frame: int = 0,
    max_frames: int = 0,
    stride: int = 1,
    save_snapshots: bool = True,
    save_fastad_overlay: bool = False,
    category: str = "DEFAULT",
) -> Dict[str, Any]:
    """单视频离线评测"""
    video_name = os.path.basename(video_path)
    cap = cv2.VideoCapture(video_path)
    if not cap.isOpened():
        print(f"[ERROR] 无法打开视频: {video_path}")
        return {"error": "cannot open video", "video": video_name, "video_path": video_path}

    if start_frame > 0:
        cap.set(cv2.CAP_PROP_POS_FRAMES, start_frame)

    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    fps = cap.get(cv2.CAP_PROP_FPS) or 25.0
    width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
    height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))

    target_frames = total_frames if max_frames <= 0 else min(total_frames, max_frames)
    print(f"\n========================================================")
    print(f"▶ 正在测试视频: {video_name} (类别: {category})")
    print(f"  分辨率: {width}x{height} | 帧率: {fps:.1f} FPS | 目标总帧: {target_frames} (采样步长: {stride})")
    print(f"========================================================")

    # 重置调度器所有插件的状态（时序窗口、EMA历史等）及融合决策层缓存
    scheduler.reset_all()
    fusion.reset()

    save_subdir = os.path.join(output_dir, os.path.splitext(video_name)[0])
    if save_snapshots:
        os.makedirs(save_subdir, exist_ok=True)

    frame_idx = start_frame
    if start_frame > 0:
        cap.set(cv2.CAP_PROP_POS_FRAMES, start_frame)
    processed_count = 0
    latencies = []
    fault_counts: Dict[str, int] = {k: 0 for k in FAULT_NAME_MAP.keys()}
    first_trigger_frame: Dict[str, int] = {}
    alarm_events: List[Dict[str, Any]] = []

    last_active_set = set()

    t_start = time.time()
    try:
        while cap.isOpened():
            ret, frame = cap.read()
            if not ret or frame is None:
                break
            frame_idx += 1
            if max_frames > 0 and processed_count >= max_frames:
                break

            if stride > 1 and (frame_idx % stride != 0):
                continue

            processed_count += 1
            pts = (frame_idx - 1) / fps
            packet = FramePacket(
                frame=frame,
                timestamp=pts,
                frame_idx=frame_idx,
                car_number=6,
                camera_id=0,
                is_panto_raised=True,
                speed=0.0,
                train_id="CR400BF"
            )

            # 1. 调度层沙箱安全分发 (注入 DetectionContext 保证黑板共享与推理缓存机制正常运转)
            context = DetectionContext(packet=packet)
            t0 = time.time()
            raw_results = scheduler.dispatch(packet, context=context)
            
            # 2. 融合层综合决策
            fusion_res = fusion.decide(raw_results, packet)
            t_cost = (time.time() - t0) * 1000.0
            latencies.append(t_cost)

            # 3. 统计生效告警
            current_active = {a["fault_type"] for a in fusion_res.active_alarms}
            new_alarms = current_active - last_active_set

            if new_alarms:
                for f_type in new_alarms:
                    cn_name = FAULT_NAME_MAP.get(f_type, f_type)
                    fault_counts[f_type] = fault_counts.get(f_type, 0) + 1
                    if f_type not in first_trigger_frame:
                        first_trigger_frame[f_type] = frame_idx

                    # 记录事件
                    alarm_events.append({
                        "frame": frame_idx,
                        "timestamp_s": round(pts, 2),
                        "fault_type": f_type,
                        "chinese_name": cn_name,
                        "score": round(fusion_res.overall_score, 3),
                        "level": fusion_res.overall_level,
                    })

                    print(f"  [Frame {frame_idx:05d} / {pts:6.2f}s] [{t_cost:5.1f}ms] 🚨 触发新告警: 【{cn_name}】(score={fusion_res.overall_score:.2f})")

                    if save_snapshots:
                        hud_img = visualizer.render_hud(frame, fusion_res, packet)
                        snap_name = f"alarm_f{frame_idx:05d}_{pts:06.1f}s_{f_type}.jpg"
                        snap_path = os.path.join(save_subdir, snap_name)
                        cv2.imwrite(snap_path, hud_img)

                    # [仅测试启用] 若为 FastAD 检测的故障类型且启用了 --save-fastad-overlay，保存透明叠加异常热图
                    if save_fastad_overlay and f_type in FASTAD_FAULT_TYPES:
                        fastad_result = None
                        for key, val in context.shared_data.items():
                            if key.startswith("fastad_result:") and isinstance(val, dict) and "anomaly_map" in val:
                                fastad_result = val
                                break

                        if fastad_result is not None:
                            overlay_dir = os.path.join(save_subdir, "fastad_overlays")
                            os.makedirs(overlay_dir, exist_ok=True)
                            overlay_img = _render_fastad_diag_img(
                                frame=frame,
                                fastad_result=fastad_result,
                                fault_type=f_type,
                                score=fusion_res.overall_score,
                            )
                            overlay_path = os.path.join(
                                overlay_dir, f"fastad_f{frame_idx:05d}_{pts:06.1f}s_{f_type}.jpg"
                            )
                            cv2.imwrite(overlay_path, overlay_img)
                            print(f"    💾 [FastAD诊断热图已保存] {os.path.basename(overlay_path)}")

            last_active_set = set(current_active)

            if processed_count % 30 == 0 or processed_count == 1:
                active_str = ", ".join([FAULT_NAME_MAP.get(a, a) for a in current_active]) if current_active else "正常"
                print(f"  进度: [{frame_idx:05d}/{target_frames}] 时延: {t_cost:5.1f}ms | 状态: {active_str}")

    except KeyboardInterrupt:
        print("\n[WARN] 用户手动中断视频评测")
    finally:
        cap.release()

    total_time = time.time() - t_start
    avg_latency = float(np.mean(latencies)) if latencies else 0.0
    p95_latency = float(np.percentile(latencies, 95)) if latencies else 0.0
    fps_measured = processed_count / total_time if total_time > 0 else 0.0

    print(f"\n--- 视频 {video_name} 评测结果汇总 ---")
    print(f"处理帧数: {processed_count} 帧 | 耗时: {total_time:.2f}s | 实测平均帧率: {fps_measured:.1f} FPS")
    print(f"平均时延: {avg_latency:.2f}ms | P95 时延: {p95_latency:.2f}ms")
    print(f"告警统计: {', '.join([f'{FAULT_NAME_MAP.get(k, k)}={v}次' for k, v in fault_counts.items() if v > 0]) or '全段正常无告警'}")

    return {
        "video": video_name,
        "video_path": video_path,
        "category": category,
        "total_frames": total_frames,
        "processed_frames": processed_count,
        "avg_latency_ms": round(avg_latency, 2),
        "p95_latency_ms": round(p95_latency, 2),
        "fps": round(fps_measured, 1),
        "fault_counts": fault_counts,
        "has_alarm": bool(len(alarm_events) > 0),
        "first_trigger": first_trigger_frame,
        "alarm_events": alarm_events,
    }


def evaluate_image(
    image_path: str,
    scheduler: DetectionScheduler,
    fusion: FusionEngine,
    visualizer: Visualizer,
    output_dir: str,
    save_snapshot: bool = True,
    save_fastad_overlay: bool = False,
) -> Dict[str, Any]:
    """单张图片静态离线评测与诊断"""
    img_name = os.path.basename(image_path)
    img = cv2.imread(image_path)
    if img is None:
        print(f"[ERROR] 无法读取图片文件: {image_path}")
        return {"error": "cannot open image"}

    h, w = img.shape[:2]
    print(f"\n========================================================")
    print(f"▶ 正在测试单张图片: {img_name} (分辨率: {w}x{h})")
    print(f"========================================================")

    # 重置调度器所有插件的状态及融合层状态
    scheduler.reset_all()
    fusion.reset()

    packet = FramePacket(
        frame=img,
        timestamp=0.0,
        frame_idx=1,
        car_number=6,
        camera_id=0,
        is_panto_raised=True,
        speed=0.0,
        train_id="CR400BF"
    )
    context = DetectionContext(packet=packet, shared_data={"single_image_mode": True})

    t0 = time.time()
    # 强制调度运行所有启用的检测器插件 (force_run_all=True)
    raw_results = scheduler.dispatch(packet, context=context, force_run_all=True)
    fusion_res = fusion.decide(raw_results, packet)
    latency_ms = (time.time() - t0) * 1000.0

    print(f"【推理耗时】: {latency_ms:.1f} ms")
    print(f"【综合诊断】: {fusion_res.primary_fault} (全局存在异常: {fusion_res.has_fault})")
    print("-" * 75)
    print(f"{'检测器插件':<22} | {'是否异常':<8} | {'置信度/分数':<10} | {'告警分级':<8} | {'详细定位/诊断信息'}")
    print("-" * 75)

    fault_counts = {k: 0 for k in FAULT_NAME_MAP.keys()}
    for name, res in fusion_res.detector_results.items():
        cn_name = FAULT_NAME_MAP.get(res.fault_type, res.fault_type)
        if res.is_abnormal:
            abn_str = "🚨 异常"
            fault_counts[res.fault_type] = fault_counts.get(res.fault_type, 0) + 1
        elif res.extra.get("suppressed_by"):
            abn_str = "🛡️ 抑制"
        else:
            abn_str = "正常"

        extra_info = []
        if res.bbox and res.is_abnormal:
            extra_info.append(f"bbox={res.bbox}")
        if res.extra.get("suppressed_by"):
            extra_info.append(f"suppressed_by={res.extra['suppressed_by']}")
        if res.fault_type == "tilt":
            extra_info.append(f"rel_tilt={res.extra.get('relative_tilt_deg', 0)}°({res.extra.get('status', 'NORMAL')})")
        elif res.fault_type == "dirty":
            extra_info.append(f"cells={res.extra.get('persistent_cells', 0)}, anom={res.extra.get('image_score', 0)}")
        elif res.fault_type in ("foreign_object", "structure_abnormal") and res.is_abnormal:
            cnt = res.extra.get('foreign_count', res.extra.get('defect_count', 0))
            extra_info.append(f"count={cnt}")
        detail_str = ", ".join(extra_info) if extra_info else "-"

        tag = f"{name}({cn_name})"
        print(f"{tag:<22} | {abn_str:<8} | {res.score:<10.2f} | {res.level:<8} | {detail_str}")

    snap_path = None
    if save_snapshot:
        save_dir = os.path.join(output_dir, "images")
        os.makedirs(save_dir, exist_ok=True)
        hud_frame = visualizer.render_hud(img, fusion_res, packet)
        snap_path = os.path.join(save_dir, f"eval_{img_name}")
        cv2.imwrite(snap_path, hud_frame)
        print("-" * 75)
        print(f"📸 诊断标注图已保存至: {snap_path}")

    # [仅测试启用] 单张图片 FastAD 热图透明叠加
    if save_fastad_overlay and fusion_res.has_fault:
        for f_dict in fusion_res.active_alarms:
            ft = f_dict["fault_type"]
            if ft in FASTAD_FAULT_TYPES:
                fastad_result = None
                for key, val in context.shared_data.items():
                    if key.startswith("fastad_result:") and isinstance(val, dict) and "anomaly_map" in val:
                        fastad_result = val
                        break
                if fastad_result is not None:
                    save_dir = os.path.join(output_dir, "fastad_overlays")
                    os.makedirs(save_dir, exist_ok=True)
                    overlay_img = _render_fastad_diag_img(
                        frame=img,
                        fastad_result=fastad_result,
                        fault_type=ft,
                        score=fusion_res.overall_score,
                    )
                    diag_path = os.path.join(save_dir, f"fastad_eval_{img_name}")
                    cv2.imwrite(diag_path, overlay_img)
                    print(f"📸 FastAD 热图诊断图已保存至: {diag_path}")
                break

    return {
        "image": img_name,
        "primary_fault": fusion_res.primary_fault,
        "has_fault": fusion_res.has_fault,
        "latency_ms": round(latency_ms, 2),
        "fault_counts": fault_counts,
        "active_alarms": fusion_res.active_alarms,
        "snapshot_path": snap_path,
    }


def main():
    parser = argparse.ArgumentParser(description="弓网视频与图像离线基准评测套件")
    parser.add_argument("--config", type=str, default="config/detectors.yaml", help="检测器配置文件")
    parser.add_argument("--input", type=str, default=None, help="统一输入路径 (支持单张图片、单视频、或包含图片/视频的目录)")
    parser.add_argument("--image", type=str, default=None, help="指定单张测试图片路径")
    parser.add_argument("--image-dir", type=str, default=None, help="指定测试图片目录路径")
    parser.add_argument("--video", type=str, default=None, help="指定单个测试视频路径")
    parser.add_argument("--dir", type=str, default=None, help="基准测试目录路径 (默认: 弓网典型视频)")
    parser.add_argument("--start-frame", type=int, default=0, help="视频起始测试帧号 (默认0)")
    parser.add_argument("--max-frames", type=int, default=0, help="视频最多评测帧数 (0为全部)")
    parser.add_argument("--stride", type=int, default=1, help="抽帧步长 (1为逐帧)")
    parser.add_argument("--output-dir", type=str, default="local_cache/benchmark_results", help="结果抓拍保存目录")
    parser.add_argument("--recursive", action="store_true", help="递归扫描子目录中的视频文件（适用于类别/车型/视频.MP4 三层结构）")
    parser.add_argument("--save-fastad-overlay", action="store_true", help="[测试专用] 告警帧保存 FastAD 热图透明叠加诊断图")
    parser.add_argument("--sample-per-category", type=int, default=0, help="每个场景类别抽样 N 个视频（0 为全量）")
    parser.add_argument("--manifest-json", type=str, default=None, help="输入或输出的数据集 Manifest JSON 路径")
    parser.add_argument("--export-report-csv", type=str, default=None, help="导出逐视频指标汇总 CSV 路径")
    parser.add_argument("--disable", "--disable-detectors", type=str, default=None, help="指定停用的检测器列表，逗号分隔 (如: --disable dirt,tilt)")
    parser.add_argument("--detectors", "--only", type=str, default=None, help="仅启用的检测器列表，逗号分隔 (如: --detectors spark,tilt)")
    args = parser.parse_args()

    cfg_path = os.path.join(BASE_DIR, args.config) if not os.path.isabs(args.config) else args.config
    out_dir = os.path.join(BASE_DIR, args.output_dir) if not os.path.isabs(args.output_dir) else args.output_dir
    os.makedirs(out_dir, exist_ok=True)

    print("========================================================")
    print(" 🚀 初始化弓网分层检测调度引擎")
    print(f" 配置文件: {cfg_path}")
    print(f" 结果目录: {out_dir}")
    print(f" FastAD热图透明叠加: {'已启用 (测试模式)' if args.save_fastad_overlay else '未启用'}")
    print("========================================================")

    scheduler = DetectionScheduler(cfg_path)

    # 动态应用命令行启停参数
    if args.detectors:
        target_set = set(d.strip() for d in args.detectors.split(",") if d.strip())
        for name in list(scheduler.detector_map.keys()):
            scheduler.set_enabled(name, name in target_set)
    elif args.disable:
        disabled_set = set(d.strip() for d in args.disable.split(",") if d.strip())
        for name in disabled_set:
            if not scheduler.set_enabled(name, False):
                print(f"[WARN] 未找到指定停用的检测器: {name} (当前可用: {list(scheduler.detector_map.keys())})")

    active_dets = [d.name for d in scheduler.detectors if d.enabled]
    inactive_dets = [d.name for d in scheduler.detectors if not d.enabled]
    print(f" 当前启用检测器 ({len(active_dets)}): {', '.join(active_dets) if active_dets else '全部已禁用'}")
    if inactive_dets:
        print(f" 当前停用检测器 ({len(inactive_dets)}): {', '.join(inactive_dets)}")

    fusion = FusionEngine()
    visualizer = Visualizer()

    # 1. 判断是否为单张图片评测
    single_img_path = None
    if args.image:
        single_img_path = args.image
    elif args.input and os.path.isfile(args.input) and os.path.splitext(args.input)[1].lower() in IMAGE_EXTS:
        single_img_path = args.input
    elif args.video and os.path.isfile(args.video) and os.path.splitext(args.video)[1].lower() in IMAGE_EXTS:
        single_img_path = args.video

    if single_img_path:
        full_path = os.path.join(BASE_DIR, single_img_path) if not os.path.isabs(single_img_path) else single_img_path
        evaluate_image(
            image_path=full_path,
            scheduler=scheduler,
            fusion=fusion,
            visualizer=visualizer,
            output_dir=out_dir,
            save_fastad_overlay=args.save_fastad_overlay,
        )
        return

    # 2. 判断是否为图片目录批量评测
    img_dir_path = args.image_dir
    if not img_dir_path and args.input and os.path.isdir(args.input):
        imgs_in_dir = [f for f in os.listdir(args.input) if os.path.splitext(f)[1].lower() in IMAGE_EXTS]
        if imgs_in_dir:
            img_dir_path = args.input

    if img_dir_path:
        full_dir = os.path.join(BASE_DIR, img_dir_path) if not os.path.isabs(img_dir_path) else img_dir_path
        img_list = sorted([
            os.path.join(full_dir, f) for f in os.listdir(full_dir)
            if os.path.splitext(f)[1].lower() in IMAGE_EXTS
        ])
        print(f"\n📁 正在批量评测图片目录: {full_dir} (共 {len(img_list)} 张图像)")
        img_results = []
        for ipath in img_list:
            res = evaluate_image(
                image_path=ipath,
                scheduler=scheduler,
                fusion=fusion,
                visualizer=visualizer,
                output_dir=out_dir,
                save_fastad_overlay=args.save_fastad_overlay,
            )
            img_results.append(res)

        print("\n" + "=" * 80)
        print(" 📊 图片批量离线评测最终汇总大表")
        print("=" * 80)
        print(f"{'图片文件名':<25} | {'耗时(ms)':<10} | {'综合主诊断':<18} | {'检出异常项'}")
        print("-" * 80)
        for r in img_results:
            if "error" in r:
                continue
            iname = r["image"]
            l_ms = r["latency_ms"]
            p_fault = r["primary_fault"]
            active = [f"{FAULT_NAME_MAP.get(k, k)}:{v}" for k, v in r["fault_counts"].items() if v > 0]
            active_str = ", ".join(active) if active else "无异常"
            print(f"{iname:<25} | {l_ms:<10.1f} | {p_fault:<18} | {active_str}")
        print("=" * 80 + "\n")
        return

    # 3. 视频评测逻辑
    vdir = None
    if args.video:
        video_list = [args.video]
    elif args.input and os.path.isfile(args.input):
        video_list = [args.input]
    else:
        vdir_name = args.dir or "弓网典型视频"
        vdir = os.path.join(BASE_DIR, vdir_name) if not os.path.isabs(vdir_name) else vdir_name
        if args.recursive:
            video_list = collect_videos_recursive(vdir)
            print(f"[INFO] 递归扫描找到 {len(video_list)} 个视频文件")
        else:
            video_list = sorted([
                str(p) for p in pathlib.Path(vdir).iterdir()
                if p.is_file() and p.suffix.lower() in VIDEO_EXTS
            ])

    # 抽样处理 (如果指定 --sample-per-category)
    if args.sample_per_category > 0 and len(video_list) > 1:
        cat_map = defaultdict(list)
        for vp in video_list:
            p = pathlib.Path(vp)
            if vdir and p.is_relative_to(pathlib.Path(vdir)):
                rel_parts = p.relative_to(pathlib.Path(vdir)).parts
                cat = rel_parts[0] if len(rel_parts) > 1 else "DEFAULT"
            else:
                cat = p.parent.name
            cat_map[cat].append(vp)

        sampled_list = []
        print(f"\n--- 执行类别抽样 (每类最多 {args.sample_per_category} 个) ---")
        for cat, vids in sorted(cat_map.items()):
            selected = vids[:args.sample_per_category]
            sampled_list.extend(selected)
            print(f"  类别 【{cat}】: 现有 {len(vids)} 个 → 抽样 {len(selected)} 个")
        video_list = sampled_list
        print(f"[INFO] 抽样完成，本次评测视频总数: {len(video_list)}\n")

    all_results = []
    for vpath in video_list:
        if vpath.endswith(".bak"):
            continue

        # 解析所属类别
        p = pathlib.Path(vpath)
        if vdir and p.is_relative_to(pathlib.Path(vdir)):
            rel_parts = p.relative_to(pathlib.Path(vdir)).parts
            category = rel_parts[0] if len(rel_parts) > 1 else "DEFAULT"
        else:
            category = p.parent.name

        res = evaluate_video(
            video_path=vpath,
            scheduler=scheduler,
            fusion=fusion,
            visualizer=visualizer,
            output_dir=out_dir,
            start_frame=args.start_frame,
            max_frames=args.max_frames,
            stride=args.stride,
            save_snapshots=True,
            save_fastad_overlay=args.save_fastad_overlay,
            category=category,
        )
        all_results.append(res)

    # ── 4. 统计汇总大表输出 ───────────────────────────────────
    valid_results = [r for r in all_results if "error" not in r]
    total_videos = len(valid_results)
    alarm_videos = sum(1 for r in valid_results if r.get("has_alarm", False))
    suppressed_videos = total_videos - alarm_videos
    fasr = (suppressed_videos / total_videos * 100.0) if total_videos > 0 else 0.0

    print("\n" + "=" * 90)
    print(" 📊 全场景基准离线评测逐视频汇总明细")
    print("=" * 90)
    print(f"{'视频文件名':<22} | {'场景类别':<15} | {'帧数':<6} | {'FPS':<6} | {'均延(ms)':<8} | {'告警判定':<8} | {'检出故障'}")
    print("-" * 90)
    for r in valid_results:
        vname = r["video"]
        cat = r.get("category", "DEFAULT")
        n_f = r["processed_frames"]
        m_fps = r["fps"]
        avg_l = r["avg_latency_ms"]
        status_str = "🚨 有告警" if r.get("has_alarm") else "🛡️ 正常(无误报)"
        active = [f"{FAULT_NAME_MAP.get(k, k)}:{v}" for k, v in r["fault_counts"].items() if v > 0]
        active_str = ", ".join(active) if active else "-"
        print(f"{vname:<22} | {cat:<15} | {n_f:<6} | {m_fps:<6.1f} | {avg_l:<8.1f} | {status_str:<8} | {active_str}")
    print("=" * 90)

    # 分类别统计
    cat_summary = defaultdict(lambda: {"total": 0, "alarms": 0, "latencies": [], "fps": []})
    for r in valid_results:
        cat = r.get("category", "DEFAULT")
        cat_summary[cat]["total"] += 1
        if r.get("has_alarm"):
            cat_summary[cat]["alarms"] += 1
        cat_summary[cat]["latencies"].append(r["avg_latency_ms"])
        cat_summary[cat]["fps"].append(r["fps"])

    print("\n" + "=" * 80)
    print(" 📈 分场景类别误报抑制率与性能收益 (FASR Breakdown)")
    print("=" * 80)
    print(f"{'场景类别':<22} | {'评测数':<6} | {'告警数':<6} | {'无误报数':<6} | {'误报抑制率(FASR)':<16} | {'平均FPS'}")
    print("-" * 80)
    for cat, info in sorted(cat_summary.items()):
        c_tot = info["total"]
        c_alr = info["alarms"]
        c_sup = c_tot - c_alr
        c_fasr = (c_sup / c_tot * 100.0) if c_tot > 0 else 0.0
        c_fps = np.mean(info["fps"]) if info["fps"] else 0.0
        print(f"{cat:<22} | {c_tot:<6} | {c_alr:<6} | {c_sup:<6} | {c_fasr:>14.1f}% | {c_fps:>6.1f}")
    print("-" * 80)
    print(f"【总体统计】 评测视频: {total_videos} 个 | 告警触发: {alarm_videos} 个 | 成功抑制虚警: {suppressed_videos} 个")
    print(f"【核心指标】 总体误报抑制率 (Overall FASR): {fasr:.2f}%")
    print("=" * 80 + "\n")

    # 导出逐视频 CSV
    if args.export_report_csv:
        csv_path = pathlib.Path(args.export_report_csv).resolve()
        csv_path.parent.mkdir(parents=True, exist_ok=True)
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            fieldnames = [
                "video", "category", "video_path", "total_frames",
                "processed_frames", "fps", "avg_latency_ms", "p95_latency_ms",
                "has_alarm", "alarm_faults", "alarm_count"
            ]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for r in valid_results:
                alr_faults = [FAULT_NAME_MAP.get(k, k) for k, v in r["fault_counts"].items() if v > 0]
                alr_cnt = sum(r["fault_counts"].values())
                writer.writerow({
                    "video": r["video"],
                    "category": r.get("category", "DEFAULT"),
                    "video_path": r.get("video_path", ""),
                    "total_frames": r["total_frames"],
                    "processed_frames": r["processed_frames"],
                    "fps": r["fps"],
                    "avg_latency_ms": r["avg_latency_ms"],
                    "p95_latency_ms": r["p95_latency_ms"],
                    "has_alarm": r.get("has_alarm", False),
                    "alarm_faults": ";".join(alr_faults),
                    "alarm_count": alr_cnt,
                })
        print(f"📄 逐视频离线评测报表已导出至: {csv_path}")

    # 回填 Manifest JSON (若指定)
    if args.manifest_json and os.path.isfile(args.manifest_json):
        try:
            with open(args.manifest_json, "r", encoding="utf-8") as f:
                manifest_records = json.load(f)

            res_map = {r["video_path"]: r for r in valid_results}
            res_map_name = {r["video"]: r for r in valid_results}

            updated_cnt = 0
            for item in manifest_records:
                match_r = res_map.get(item.get("file_path")) or res_map_name.get(item.get("filename"))
                if match_r:
                    item["our_alarm"] = match_r.get("has_alarm", False)
                    item["our_fault_types"] = [
                        FAULT_NAME_MAP.get(k, k) for k, v in match_r["fault_counts"].items() if v > 0
                    ]
                    updated_cnt += 1

            with open(args.manifest_json, "w", encoding="utf-8") as f:
                json.dump(manifest_records, f, ensure_ascii=False, indent=2)

            # 同步写回对应的 CSV 文件 (支持 Excel 直接打开查看评测回填)
            manifest_csv = pathlib.Path(args.manifest_json).with_suffix(".csv")
            if manifest_records:
                with open(manifest_csv, "w", encoding="utf-8-sig", newline="") as f:
                    writer = csv.DictWriter(f, fieldnames=list(manifest_records[0].keys()))
                    writer.writeheader()
                    for r in manifest_records:
                        r_copy = dict(r)
                        if isinstance(r_copy.get("our_fault_types"), list):
                            r_copy["our_fault_types"] = ";".join(r_copy["our_fault_types"])
                        writer.writerow(r_copy)

            # 同步写回对应的 Excel 文件 (保留下拉框与原有单元格样式)
            manifest_xlsx = pathlib.Path(args.manifest_json).with_suffix(".xlsx")
            if manifest_xlsx.exists():
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(manifest_xlsx)
                    ws = wb.active
                    header_row = [cell.value for cell in ws[1]]
                    col_map = {val: idx for idx, val in enumerate(header_row, 1) if val}

                    if "our_alarm" in col_map and "our_fault_types" in col_map:
                        fn_col = col_map.get("filename")
                        fp_col = col_map.get("file_path")
                        alarm_col = col_map["our_alarm"]
                        fault_col = col_map["our_fault_types"]

                        for row_i in range(2, ws.max_row + 1):
                            fn_val = ws.cell(row=row_i, column=fn_col).value if fn_col else None
                            fp_val = ws.cell(row=row_i, column=fp_col).value if fp_col else None
                            match_r = res_map.get(fp_val) or res_map_name.get(fn_val)
                            if match_r:
                                has_alr = match_r.get("has_alarm", False)
                                alr_faults = [
                                    FAULT_NAME_MAP.get(k, k) for k, v in match_r["fault_counts"].items() if v > 0
                                ]
                                ws.cell(row=row_i, column=alarm_col, value=str(has_alr))
                                ws.cell(row=row_i, column=fault_col, value=";".join(alr_faults))

                        wb.save(manifest_xlsx)
                except Exception as ex:
                    print(f"[WARN] 同步更新 Manifest Excel 失败: {ex}")

            print(f"📝 已同步更新 Manifest (JSON, CSV & Excel) 中 {updated_cnt} 个视频的评测状态: {args.manifest_json}")
        except Exception as e:
            print(f"[WARN] 更新 Manifest 文件失败: {e}")


if __name__ == "__main__":
    main()
